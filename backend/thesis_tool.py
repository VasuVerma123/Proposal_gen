import os
import re
import json
import logging
from io import BytesIO
from pathlib import Path

import openai
import markdown
from xhtml2pdf import pisa
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)

THESYS_API_KEY = os.environ.get(
    "THESYS_API_KEY",
    "sk-th-ABwCkZkFin6cEoEzCvH8NwkGSasr1uBEjmi66sjvyt302qJIm1DAGWx6NtjXW7QwpdyhkBzWhXGXNndWFgeL1UnMTMRaqTs6XtkE",
)
ARTIFACT_BASE_URL = "https://api.thesys.dev/v1/artifact"
ARTIFACT_MODEL = "c1/artifact/v-20260130"

HEADING_COLOR = RGBColor(0x0F, 0x34, 0x60)


# ═══════════════════════════════════════════════════════════════════════════
#  1. THESYS C1 API
# ═══════════════════════════════════════════════════════════════════════════

def call_c1_api(markdown_text: str) -> str:
    """Send markdown to the Thesys C1 Artifact API and return the raw response."""
    client = openai.OpenAI(base_url=ARTIFACT_BASE_URL, api_key=THESYS_API_KEY)

    logger.info("Sending content to Thesys C1 API ...")

    completion = client.chat.completions.create(
        model=ARTIFACT_MODEL,
        messages=[{
            "role": "user",
            "content": (
                "Convert the following content into a professional report. "
                "Preserve all sections, bullet points, tables, and details exactly.\n\n"
                + markdown_text
            ),
        }],
        metadata={
            "thesys": json.dumps({
                "c1_artifact_type": "report",
                "id": "thesis-report",
            })
        },
    )

    c1_response = completion.choices[0].message.content
    logger.info("C1 API returned %d chars", len(c1_response))
    return c1_response


# ═══════════════════════════════════════════════════════════════════════════
#  2. SAVE FULL CONTENT AS PDF
# ═══════════════════════════════════════════════════════════════════════════

def save_as_pdf(md_text: str, output_path: str) -> Path:
    """Convert the full markdown content to a styled PDF."""
    body_html = markdown.markdown(md_text, extensions=["tables", "fenced_code"])

    full_html = (
        "<!DOCTYPE html><html><head><meta charset='UTF-8'/>"
        "<style>"
        "@page { size: A4; margin: 2cm; }"
        "body { font-family: Helvetica, Arial, sans-serif; color: #1a1a2e; "
        "       line-height: 1.6; font-size: 13px; }"
        "h1 { color: #0f3460; font-size: 24px; border-bottom: 3px solid #e94560; "
        "     padding-bottom: 6px; margin-top: 28px; }"
        "h2 { color: #0f3460; font-size: 19px; border-bottom: 1px solid #ccc; "
        "     padding-bottom: 4px; margin-top: 22px; }"
        "h3 { color: #16213e; font-size: 15px; margin-top: 18px; }"
        "hr { border: none; border-top: 1px solid #ddd; margin: 20px 0; }"
        "ul { padding-left: 22px; }"
        "li { margin-bottom: 3px; }"
        "table { width: 100%; border-collapse: collapse; margin: 12px 0; }"
        "th { background: #0f3460; color: #fff; padding: 8px 10px; text-align: left; }"
        "td { padding: 8px 10px; border-bottom: 1px solid #ddd; }"
        "strong { color: #0f3460; }"
        "p { margin: 6px 0; }"
        "</style></head><body>"
        f"{body_html}"
        "</body></html>"
    )

    out = Path(output_path)
    buf = BytesIO()
    status = pisa.CreatePDF(full_html, dest=buf)
    if status.err:
        raise RuntimeError(f"PDF generation failed with {status.err} error(s)")
    out.write_bytes(buf.getvalue())
    logger.info("PDF saved: %s (%d bytes)", out.resolve(), out.stat().st_size)
    return out


# ═══════════════════════════════════════════════════════════════════════════
#  3. SAVE FULL CONTENT AS DOCX
# ═══════════════════════════════════════════════════════════════════════════

def save_as_docx(md_text: str, output_path: str) -> Path:
    """Convert the full markdown content to a styled DOCX."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:], level=1)
            for r in p.runs:
                r.font.color.rgb = HEADING_COLOR
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_heading(stripped[3:], level=2)
            for r in p.runs:
                r.font.color.rgb = HEADING_COLOR
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_heading(stripped[4:], level=3)
            for r in p.runs:
                r.font.color.rgb = HEADING_COLOR
            i += 1
            continue

        if stripped.startswith("* ") or stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^\d+\.\s", "", stripped)
            doc.add_paragraph(text, style="List Number")
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_md_table(doc, table_lines)
            continue

        p = doc.add_paragraph()
        _add_rich_text(p, stripped)
        i += 1

    out = Path(output_path)
    doc.save(str(out))
    logger.info("DOCX saved: %s (%d bytes)", out.resolve(), out.stat().st_size)
    return out


def _add_rich_text(paragraph, text: str):
    """Parse **bold** markers in text and add runs accordingly."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def _add_md_table(doc: Document, table_lines: list[str]):
    """Parse markdown table lines and add a styled table to the document."""
    rows = []
    for tl in table_lines:
        cells = [c.strip() for c in tl.strip("|").split("|")]
        if all(set(c) <= set("- :") for c in cells):
            continue  # separator row
        rows.append(cells)

    if not rows:
        return

    ncols = len(rows[0])
    tbl = doc.add_table(rows=1, cols=ncols)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ci, val in enumerate(rows[0]):
        tbl.rows[0].cells[ci].text = val
        for r in tbl.rows[0].cells[ci].paragraphs[0].runs:
            r.bold = True

    for row_data in rows[1:]:
        row = tbl.add_row()
        for ci, val in enumerate(row_data):
            if ci < ncols:
                row.cells[ci].text = val

    doc.add_paragraph("")


# ═══════════════════════════════════════════════════════════════════════════
#  4. MAIN
# ═══════════════════════════════════════════════════════════════════════════

def generate_and_download(
    markdown_text: str,
    pdf_path: str = "output.pdf",
    docx_path: str = "output.docx",
) -> tuple[Path, Path]:
    """
    1. Send markdown (LLM response) to Thesys C1 API.
    2. Save the full content as PDF and DOCX.
    """
    c1_response = call_c1_api(markdown_text)
    logger.info("C1 API response received. Saving full content as PDF and DOCX ...")

    pdf_out = save_as_pdf(markdown_text, pdf_path)
    docx_out = save_as_docx(markdown_text, docx_path)
    return pdf_out, docx_out


if __name__ == "__main__":
    md_file = Path("test.md")
    if not md_file.exists():
        logger.error("test.md not found in %s", Path.cwd())
        raise SystemExit(1)

    llm_response = md_file.read_text(encoding="utf-8")
    logger.info("Loaded LLM response from %s (%d chars)", md_file, len(llm_response))

    pdf_out, docx_out = generate_and_download(
        markdown_text=llm_response,
        pdf_path="telecom_proposal.pdf",
        docx_path="telecom_proposal.docx",
    )

    logger.info("Done!")
    logger.info("  PDF:  %s", pdf_out.resolve())
    logger.info("  DOCX: %s", docx_out.resolve())
